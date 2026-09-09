#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
KOSA 기능점수(FP) 간이법 SW 개발비 견적서 생성기.

사용법:
    python build_estimate.py spec.json [--template TEMPLATE.xlsx] [--output OUT.xlsx]

spec.json 스키마는 references/spec-schema.md 참고.
산식/보정계수는 모두 엑셀 수식으로 들어가므로, 결과 파일을 Excel/Google Sheets에서
열면 자동 재계산된다(LibreOffice 없이도 OK). 금액 미리보기는 calc_fp.py로 확인.
"""
import argparse, json, os, sys
from copy import copy
import warnings; warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.workbook.properties import CalcProperties

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "..", "assets", "kosa_fp_template.xlsx")

# 보정계수 수준(1~5) -> '(참조)보정계수' 시트 lookup 텍스트(접두 숫자로 근사매칭)
LINKAGE = {
 1:"1. 타 기관 연계 없음", 2:"2. 1~2개의 타 기관 연계", 3:"3. 3~5개의 타 기관 연계",
 4:"4. 6~10개의 타 기관 연계", 5:"5. 10개 초과의 타 기관 연계"}
PERF = {
 1:"1. 응답성능에 대한 특별한 요구사항이 없다.",
 2:"2. 응답성능에 대한 요구사항이 있으나 특별한 조치가 필요하지는 않다.",
 3:"3. 응답시간이나 처리율이 피크타임(peak time)에 중요하며, 처리 시한이 명시되어 있다.",
 4:"4. 응답시간이나 처리율이 모든 업무시간에 중요하며, 처리 시한이 명시되어 있다.",
 5:"5. 응답성능 요구수준이 엄격하여, 설계, 개발 또는 구현 단계에서 성능 분석도구 사용이 필요하다."}
OPERATION = {
 1:"1. 운영환경 호환성에 대한 요구사항이 없다.",
 2:"2. 운영환경 호환성에 대한 요구사항이 있으며, 동일 하드웨어 및 소프트웨어 환경에서 운영되도록 설계된다.",
 3:"3. 운영환경 호환성에 대한 요구사항이 있으며, 유사 하드웨어 및 소프트웨어 환경에서 운영되도록 설계된다.",
 4:"4. 운영환경 호환성에 대한 요구사항이 있으며, 이질적인 하드웨어 및 소프트웨어 환경에서 운영되도록 설계된다.",
 5:"5. 항목 4에 더하여 일반적 산출물 이외에 장소에서 원활한 운영을 보장하기 위한 운영 절차의 문서화와 사전 모의훈련이 요구된다."}
SECURITY = {
 1:"1. 암호화, 웹취약점 점검, 시큐어코딩, 개인정보보호 등 1가지 보안 요구사항이 포함되어 있다.",
 2:"2. 2가지 요구사항이 포함되어 있다.",
 3:"3. 3가지 요구사항이 포함되어 있다.",
 4:"4. 4가지 항목이 모두 포함되어 있다.",
 5:"5. 5가지 이상의 보안 요구사항이 포함되어 있다."}

VALID_TYPES = {"EI","EO","EQ","ILF","EIF"}

def weight_formula(r):
    return (f'=IF(F{r}="","",IF(F{r}="EI",4,IF(F{r}="EO",5.2,'
            f'IF(F{r}="EQ",3.9,IF(F{r}="ILF",7.5,IF(F{r}="EIF",5.4))))))')

def build(spec, template, output):
    wb = openpyxl.load_workbook(template)

    # ===== 상세기능 =====
    ws = wb["상세기능"]
    # 스타일 도너: 템플릿 5행(B~H)
    styles = {c: copy(ws.cell(row=5, column=c)._style) for c in range(2, 9)}
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 5:
            ws.unmerge_cells(str(rng))
    for r in range(5, ws.max_row + 1):
        for c in range(2, 12):
            ws.cell(row=r, column=c).value = None

    r = 5
    app_ranges, sub_ranges = [], []
    for app in spec["applications"]:
        a0 = r
        for sub in app["subs"]:
            s0 = r
            for p in sub["processes"]:
                ftype = p["type"].strip().upper()
                if ftype not in VALID_TYPES:
                    raise ValueError(f"잘못된 FP유형 '{ftype}' (행 {r}). 허용: {VALID_TYPES}")
                ws.cell(r, 2).value = app["name"]
                ws.cell(r, 3).value = sub["name"]
                ws.cell(r, 4).value = p["name"]
                ws.cell(r, 5).value = p.get("desc", "")
                ws.cell(r, 6).value = ftype
                ws.cell(r, 7).value = weight_formula(r)
                ws.cell(r, 8).value = p.get("note", "")
                for c in range(2, 9):
                    ws.cell(r, c)._style = copy(styles[c])
                r += 1
            sub_ranges.append((s0, r - 1))
        app_ranges.append((a0, r - 1))
    last = r - 1
    for s, e in app_ranges:
        if e > s: ws.merge_cells(start_row=s, start_column=2, end_row=e, end_column=2)
    for s, e in sub_ranges:
        if e > s: ws.merge_cells(start_row=s, start_column=3, end_row=e, end_column=3)

    # ===== 상세비용 =====
    wc = wb["상세비용"]
    if spec.get("project_title"):
        wc["B1"].value = spec["project_title"]
    if spec.get("unit_price"):
        wc["D7"].value = int(spec["unit_price"])
    cor = spec.get("correction", {})
    wc["D24"].value = LINKAGE[int(cor.get("linkage", 2))]
    wc["D25"].value = PERF[int(cor.get("performance", 2))]
    wc["D26"].value = OPERATION[int(cor.get("operation", 2))]
    wc["D27"].value = SECURITY[int(cor.get("security", 2))]
    wc["J9"].value = float(spec.get("profit_rate", 0.10))

    # 직접경비 (rows 15~18, 최대 4줄). 금액은 숫자 또는 '='로 시작하는 수식.
    dcs = spec.get("direct_costs", [])
    if len(dcs) > 4:
        raise ValueError("direct_costs 는 최대 4줄까지 지원합니다. 항목을 통합하세요.")
    for i in range(4):
        row = 15 + i
        if i < len(dcs):
            d = dcs[i]
            wc.cell(row, 2).value = d["item"]
            wc.cell(row, 4).value = d.get("detail", "")
            wc.cell(row, 11).value = d["amount"]
        else:
            wc.cell(row, 2).value = None; wc.cell(row, 4).value = None; wc.cell(row, 11).value = None
    wc["K19"].value = "=SUM(K15:L18)"

    # 실비 (Pass-through, rows 38~39, 최대 2줄)
    exs = spec.get("expenses", [])
    if len(exs) > 2:
        raise ValueError("expenses(실비) 는 최대 2줄까지 지원합니다. 항목을 통합하세요.")
    for i in range(2):
        row = 38 + i
        if i < len(exs):
            e = exs[i]
            wc.cell(row, 2).value = e["item"]
            wc.cell(row, 4).value = e.get("detail", "")
            wc.cell(row, 11).value = e["amount"]
        else:
            wc.cell(row, 2).value = None; wc.cell(row, 4).value = None; wc.cell(row, 11).value = None
    wc["K40"].value = "=SUM(K38:K39)"

    # COUNTIF/SUMIF 범위가 마지막 행을 포함하는지 확인 (템플릿 기본 $F$500)
    if last > 500:
        for cell in ["E31","F31","G31","H31","I31"]:
            wc[cell].value = wc[cell].value.replace("$F$500", f"$F${last+5}")
        for cell in ["E32","F32","G32","H32","I32"]:
            wc[cell].value = wc[cell].value.replace("$F$500", f"$F${last+5}").replace("$G$500", f"$G${last+5}")

    # ===== 최종견적 =====
    wf = wb["최종견적"]
    wf["K8"].value = float(spec.get("discount_rate", 0.0))
    # 실비가 있으면 최종 견적에 별도 라인으로 반영(할인 미적용, K9 절삭 이후 합산)
    if exs:
        wf["B14"].value = "실비 (Pass-through, 할인 미적용)"
        wf["K14"].value = "=상세비용!K40"
    else:
        wf["B14"].value = None
        wf["K14"].value = None

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(output)
    return last - 4  # 함수 개수

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    ap.add_argument("--output", default=None)
    a = ap.parse_args()
    with open(a.spec, encoding="utf-8") as f:
        spec = json.load(f)
    out = a.output or spec.get("output") or "견적서.xlsx"
    n = build(spec, a.template, out)
    print(f"OK: {out}  (단위프로세스 {n}개)")
    print("→ 금액 미리보기:  python calc_fp.py", a.spec)

if __name__ == "__main__":
    main()
